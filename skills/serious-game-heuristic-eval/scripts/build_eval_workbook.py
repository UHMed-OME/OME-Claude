#!/usr/bin/env python3
"""
Build a serious-game heuristic-evaluation workbook from a findings JSON file.

Usage:
    python3 build_eval_workbook.py findings.json output.xlsx

findings.json schema (all sections optional except "game"):
{
  "game": {"name": "...", "type": "...", "stage": "...", "audience": "...",
           "objective": "...", "access_notes": "..."},
  "summary": "one short paragraph",
  "strengths": ["...", "..."],
  "procedural_rhetoric": {"intended": "...", "actual": "...",
                          "aligned": "Yes|Partially|No", "note": "..."},
  "issues": [{"title": "...", "heuristic": "...", "severity": 0-4,
              "evidence": "...", "recommendation": "..."}],
  "hep":  [{"category": "...", "heuristic": "...",
            "rating": "Pass|Partial|Fail|N/A", "severity": 0-4, "note": "..."}],
  "play": [{"category": "...", "heuristic": "...",
            "rating": "Pass|Partial|Fail|N/A", "severity": 0-4, "note": "..."}],
  "sdt": [{"need": "Autonomy|Competence|Relatedness",
           "level": "High|Medium|Low", "note": "..."}],
  "assets": [{"tier": "Tier 1|Tier 2|Tier 3", "category": "...", "asset": "...",
              "spec": "...", "format": "...", "size": "...", "qty": "...",
              "addresses": "heuristic/need", "payoff": "...",
              "source": "license-friendly source + license, e.g. 'Kenney.nl (CC0)'"}],
  "sources": ["...", "..."]
}
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY="22364F"; WHITE="FFFFFF"
RATING_FILL={"Pass":"C6E7CE","Partial":"FCE4A6","Fail":"F4C7C3","N/A":"E2E5E9"}
SEV_FILL={0:"C6E7CE",1:"E2E5E9",2:"FCE4A6",3:"F6C99B",4:"F4C7C3"}
TIER_FILL={"Tier 1":"1F6B3B","Tier 2":"8A6D1F","Tier 3":"55408A"}
thin=Side(style="thin",color="C0C8D2")
BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
HFONT=Font(name="Arial",size=10,bold=True,color=WHITE)
TFONT=Font(name="Arial",size=14,bold=True,color="1F4E79")
CFONT=Font(name="Arial",size=10,color="1A1A1A")
WRAP=Alignment(wrap_text=True,vertical="top")
CTR=Alignment(horizontal="center",vertical="center",wrap_text=True)
HFILL=PatternFill("solid",fgColor=NAVY)

def header(ws,row,cols):
    for i,_ in enumerate(cols,1):
        c=ws.cell(row=row,column=i); c.font=HFONT; c.fill=HFILL; c.alignment=CTR; c.border=BORDER

def widths(ws,ws_widths):
    for i,w in enumerate(ws_widths): ws.column_dimensions[chr(65+i)].width=w

def body(ws,r0,r1,ncols,h=42):
    for r in range(r0,r1+1):
        for c in range(1,ncols+1):
            cell=ws.cell(row=r,column=c); cell.font=CFONT; cell.alignment=WRAP; cell.border=BORDER
        ws.row_dimensions[r].height=h

def checklist_sheet(wb,title,rows):
    ws=wb.create_sheet(title)
    cols=["Category","Heuristic","Rating","Severity","Note"]
    ws.append(cols); header(ws,1,cols)
    for r in rows:
        ws.append([r.get("category",""),r.get("heuristic",""),r.get("rating",""),
                   r.get("severity",""),r.get("note","")])
    body(ws,2,1+len(rows),len(cols))
    for i,r in enumerate(rows,2):
        rc=ws.cell(row=i,column=3); rc.alignment=CTR
        if r.get("rating") in RATING_FILL: rc.fill=PatternFill("solid",fgColor=RATING_FILL[r["rating"]])
        sc=ws.cell(row=i,column=4); sc.alignment=CTR
        sev=r.get("severity")
        if isinstance(sev,int) and sev in SEV_FILL: sc.fill=PatternFill("solid",fgColor=SEV_FILL[sev])
    widths(ws,[22,40,12,12,55]); ws.freeze_panes="A2"
    return ws

def main():
    if len(sys.argv)<3:
        print("Usage: python3 build_eval_workbook.py findings.json output.xlsx"); sys.exit(1)
    data=json.load(open(sys.argv[1])); out=sys.argv[2]
    wb=Workbook(); wb.remove(wb.active)

    # Summary
    ws=wb.create_sheet("Summary")
    g=data.get("game",{})
    ws["A1"]=f"Heuristic Evaluation — {g.get('name','(game)')}"; ws["A1"].font=TFONT
    ws.merge_cells("A1:B1")
    meta=[("Type",g.get("type","")),("Stage",g.get("stage","")),("Audience",g.get("audience","")),
          ("Learning objective",g.get("objective","")),("Access / scope",g.get("access_notes",""))]
    r=3
    for k,v in meta:
        ws.cell(row=r,column=1,value=k).font=Font(name="Arial",bold=True,size=10)
        c=ws.cell(row=r,column=2,value=v); c.alignment=WRAP; r+=1
    r+=1
    ws.cell(row=r,column=1,value="Summary").font=Font(name="Arial",bold=True,size=11); r+=1
    c=ws.cell(row=r,column=1,value=data.get("summary","")); c.alignment=WRAP; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2); ws.row_dimensions[r].height=80; r+=2
    pr=data.get("procedural_rhetoric")
    if pr:
        ws.cell(row=r,column=1,value="Procedural-rhetoric verdict").font=Font(name="Arial",bold=True,size=11); r+=1
        for k in ["intended","actual","aligned","note"]:
            if pr.get(k):
                ws.cell(row=r,column=1,value=k.capitalize()).font=Font(name="Arial",bold=True,size=10)
                cc=ws.cell(row=r,column=2,value=pr.get(k)); cc.alignment=WRAP; r+=1
    if data.get("strengths"):
        r+=1; ws.cell(row=r,column=1,value="Strengths").font=Font(name="Arial",bold=True,size=11); r+=1
        for s in data["strengths"]:
            ws.cell(row=r,column=1,value="• "+s).alignment=WRAP; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2); r+=1
    widths(ws,[22,75])

    # Issues
    iss=data.get("issues",[])
    if iss:
        ws=wb.create_sheet("Issues")
        cols=["Severity","Title","Heuristic","Evidence","Recommendation"]
        ws.append(cols); header(ws,1,cols)
        for it in sorted(iss,key=lambda x:-(x.get("severity") or 0)):
            ws.append([it.get("severity",""),it.get("title",""),it.get("heuristic",""),
                       it.get("evidence",""),it.get("recommendation","")])
        body(ws,2,1+len(iss),len(cols),h=60)
        for i,it in enumerate(sorted(iss,key=lambda x:-(x.get("severity") or 0)),2):
            sc=ws.cell(row=i,column=1); sc.alignment=CTR
            sev=it.get("severity")
            if isinstance(sev,int) and sev in SEV_FILL: sc.fill=PatternFill("solid",fgColor=SEV_FILL[sev])
        widths(ws,[10,30,28,45,45]); ws.freeze_panes="A2"

    if data.get("play"): checklist_sheet(wb,"PLAY Checklist",data["play"])
    if data.get("hep"):  checklist_sheet(wb,"HEP Checklist",data["hep"])

    # SDT
    sdt=data.get("sdt",[])
    if sdt:
        ws=wb.create_sheet("Motivation (SDT)")
        cols=["Need","Level","Note"]; ws.append(cols); header(ws,1,cols)
        for s in sdt: ws.append([s.get("need",""),s.get("level",""),s.get("note","")])
        body(ws,2,1+len(sdt),len(cols),h=50); widths(ws,[18,14,70]); ws.freeze_panes="A2"

    # Assets
    assets=data.get("assets",[])
    if assets:
        ws=wb.create_sheet("Asset Recommendations")
        cols=["Tier","Category","Asset","Spec / Description","Format","Size","Qty","Addresses","Payoff","Source (license-friendly)"]
        ws.append(cols); header(ws,1,cols)
        order={"Tier 1":0,"Tier 2":1,"Tier 3":2}
        for a in sorted(assets,key=lambda x:order.get(x.get("tier",""),9)):
            ws.append([a.get("tier",""),a.get("category",""),a.get("asset",""),a.get("spec",""),
                       a.get("format",""),a.get("size",""),a.get("qty",""),a.get("addresses",""),a.get("payoff",""),a.get("source","")])
        body(ws,2,1+len(assets),len(cols),h=46)
        for i,a in enumerate(sorted(assets,key=lambda x:order.get(x.get("tier",""),9)),2):
            tc=ws.cell(row=i,column=1); tc.alignment=CTR
            if a.get("tier") in TIER_FILL:
                tc.fill=PatternFill("solid",fgColor=TIER_FILL[a["tier"]]); tc.font=Font(name="Arial",bold=True,size=10,color=WHITE)
        widths(ws,[9,12,20,30,11,10,6,18,24,30]); ws.freeze_panes="A2"

    # Sources
    src=data.get("sources",[])
    if src:
        ws=wb.create_sheet("Sources")
        ws.append(["Source"]); header(ws,1,["Source"])
        for s in src: ws.append([s])
        body(ws,2,1+len(src),1,h=30); widths(ws,[110])

    wb.save(out); print(f"Wrote {out} with sheets: {wb.sheetnames}")

if __name__=="__main__":
    main()
